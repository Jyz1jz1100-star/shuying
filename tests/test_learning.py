from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from videosummarizer.checkpoints import atomic_json
from videosummarizer.config import Settings
from videosummarizer.database import Database
from videosummarizer.evidence import make_workspace, apply_edit, normalize_block, citation, render_markdown
from videosummarizer.learning import LearningPipeline, CitedBlock
from videosummarizer.manager import JobManager
from videosummarizer.llm_settings import LLMRuntimeSettings
from videosummarizer.schemas import Transcript, TranscriptSegment
from videosummarizer.workspace_api import router_for


def raw_transcript(n=1):
    return {'language':'zh', 'source':'human_subtitles', 'segments':[
        {'start': i*3, 'end':i*3+2, 'text':f'第{i}段：GPU 不需要 32 GB。'} for i in range(n)]}


@pytest.fixture
def project(tmp_path):
    config = Settings(data_dir=tmp_path, reading_products=False)
    config.ensure_directories()
    db = Database(config.db_path)
    db.initialize()
    settings = SimpleNamespace(runtime=lambda _: LLMRuntimeSettings('local','fake-model'))
    pipeline = LearningPipeline(config,db,settings)
    manager = JobManager(db,pipeline)
    app = FastAPI()
    app.include_router(router_for(db,pipeline,manager,SimpleNamespace(active=False),config))
    return config,db,pipeline,manager,TestClient(app)


def create(project, n=1):
    config, db, pipeline, manager, client = project
    directory = config.jobs_dir/'test'
    directory.mkdir()
    db.create_job('test','',directory)
    db.update_job('test',status='failed',title='测试讲义')
    atomic_json(directory/'transcript_raw.json',raw_transcript(n))
    return directory


def test_desktop_products_edit_regenerate_and_export(project, monkeypatch):
    from dataclasses import replace
    from videosummarizer.reading_products import SummarySection, MapSection
    create(project)
    config, db, pipeline, _, client = project
    assert Settings().reading_products is True
    pipeline.config = replace(config, reading_products=True)
    calls = []

    class Writer:
        def __init__(self, *args): pass
        def ensure_model(self, *args): pass
        def unload_model(self): pass
        def _chat_model(self, schema, prompt):
            calls.append(schema)
            assert '术语表：' in prompt
            if schema is SummarySection:
                return schema(heading='硬件需求', takeaway='核对显存需求', points=[{'text':'无需大显存', 'segment_ids':['s000001']}])
            assert schema is MapSection
            return schema(topic='硬件', branches=[{'label':'显存', 'children':[{'label':'需求核对', 'segment_ids':['s000001']}]}])

    monkeypatch.setattr('videosummarizer.learning.OllamaSummarizer', Writer)
    pipeline.run('test')
    result = client.get('/api/jobs/test/workspace').json()
    assert result['summary'][0]['takeaway'] == '核对显存需求'
    assert result['mindmap'][0]['topic'] == '硬件'
    assert result['segments'][0]['text'] == raw_transcript()['segments'][0]['text']
    assert '需求核对' in client.get('/api/jobs/test/export?format=outline').text
    changed = client.patch('/api/jobs/test/segments/s000001', json={'revision':result['revision'], 'text':'新的显存要求'}).json()
    assert changed['summary'] == [] and changed['mindmap'] == []
    assert '原文已修改' in client.get('/api/jobs/test/export?format=outline').text
    pipeline.run('test')
    assert len(calls) == 4
    revision = pipeline.state(db.get_job('test'))['revision']
    changed = client.put('/api/jobs/test/glossary', json={'revision':revision, 'terms':['GPU = 图形处理器']}).json()
    assert changed['summary'] == []
    pipeline.run('test')
    assert len(calls) == 6
    pipeline.run('test')
    assert len(calls) == 6  # Resume reuses both completed products, without another model call.


def test_correction_direction_and_rejection():
    original = raw_transcript()
    corrected = raw_transcript()
    corrected['segments'][0]['text'] = 'GPU 需要 64 GB。'
    state = make_workspace(corrected,original)
    s = state['segments'][0]
    assert s['original'] == original['segments'][0]['text']
    assert s['text'] == corrected['segments'][0]['text']
    assert {'digits','negation'} <= set(s['flags'])
    reverted = apply_edit(state,s['id'],None,'reject')
    assert reverted['segments'][0]['text'] == s['original']
    assert state['segments'][0]['text'] != s['original']
    edited = apply_edit(state,s['id'],'人工修订内容','accept')
    assert edited['segments'][0]['text'] == '人工修订内容'


def test_invalid_reference_and_markdown_injection():
    state = make_workspace(raw_transcript())
    block = normalize_block({'heading':'<script>', 'paragraphs':[{'text':'![x](https://evil.test/pixel)', 'segment_ids':['unknown']}]},state['segments'],'b0')
    assert block['paragraphs'][0]['review'] == 'unverified'
    assert not block['paragraphs'][0]['segment_ids']
    state['blocks'] = [block]
    md = render_markdown(state,'<script>')
    assert '<script>' not in md and '![x](' not in md
    for url in ['javascript:alert(1)', 'http://localhost:8765', 'https://youtube.com.evil.test/watch?v=abcdefghijk']:
        assert citation(['s000001'],state['segments'],url)['url'] == ''


def test_import_without_whisper_or_llm(project):
    client = project[-1]
    content = '1\n00:00:00,000 --> 00:00:02,000\n测试字幕\n'
    response = client.post('/api/import?filename=lesson.srt',content=content.encode())
    assert response.status_code == 202
    assert response.json()['input_type'] == 'srt'
    assert client.post('/api/import?filename=bad.srt',content=b'not subtitles').status_code == 400
    assert client.post('/api/import?filename=bad.exe',content=b'MZ').status_code == 400
    invalid = '1\n00:00:05,000 --> 00:00:02,000\n倒序\n'
    assert client.post('/api/import?filename=bad.srt',content=invalid.encode()).status_code == 400


def test_navigation_snapshot_and_outline_endpoint(project):
    create(project)
    _,db,pipeline,_,client = project
    # Workspace URLs are computed for presentation, never stored as edits.
    job = db.get_job('test')
    with db.connect() as conn:
        conn.execute('UPDATE jobs SET url=? WHERE id=?', ('https://youtu.be/abcdefghijk', 'test'))
    state = make_workspace(raw_transcript())
    state['blocks'] = [normalize_block({'heading': '主题', 'paragraphs': [{'text': '笔记内容', 'segment_ids': ['s000001']}]}, state['segments'], 'b0')]
    pipeline.save(job,state)
    response = client.get('/api/jobs/test/workspace')
    assert response.status_code == 200
    assert response.json()['segments'][0]['source_url'].endswith('t=0s')
    assert 'source_url' not in pipeline.state(db.get_job('test'))['segments'][0]
    exported = client.get('/api/jobs/test/export?format=outline')
    assert exported.status_code == 200
    assert exported.headers['content-type'].startswith('text/markdown')
    assert '导图尚未生成' in exported.text and '笔记内容' not in exported.text


def test_library_does_not_hide_older_material(project):
    config,db,_,_,_ = project
    for index in range(105):
        db.create_job(f'item-{index:03}', '', config.jobs_dir/f'item-{index}')
    jobs = db.list_jobs()
    assert len(jobs) == 105
    assert {j['id'] for j in jobs} == {f'item-{i:03}' for i in range(105)}


def test_revision_stale_export_and_media(project):
    directory = create(project)
    config,db,pipeline,manager,client=project
    state = make_workspace(raw_transcript())
    state['blocks'] = [normalize_block({'heading':'章节','paragraphs':[{'text':'内容', 'segment_ids':['s000001']}]},state['segments'],'b0000')]
    pipeline.save(db.get_job('test'),state)
    response=client.patch('/api/jobs/test/segments/s000001',json={'revision':1,'text':'人工修订'})
    assert response.status_code==200
    assert response.json()['blocks'][0]['stale'] is True
    assert response.json()['segments'][0]['human_edited'] is True
    assert client.patch('/api/jobs/test/segments/s000001',json={'revision':1,'text':'覆盖'}).status_code==409
    assert (directory/'history/revision-1.json').is_file()
    for format in ['md','json','docx','srt']:
        assert client.get('/api/jobs/test/export',params={'format':format}).status_code==200
    (directory/'input.mp3').write_bytes(b'fake')
    assert client.delete('/api/jobs/test/media').status_code==200
    assert not (directory/'input.mp3').exists()
    assert (directory/'workspace.json').exists()
    assert client.get('/api/jobs/test/media').status_code==404


def test_resume_and_local_regeneration(project, monkeypatch):
    directory=create(project,161)
    config,db,pipeline,manager,client=project
    calls=[]
    fail_once=[True]
    class FakeWriter:
        def __init__(self,*args): pass
        def ensure_model(self,*args): pass
        def unload_model(self): pass
        def _chat_model(self,cls,prompt):
            data=json.loads(prompt.split('\n字幕：',1)[1].split('\n再次核对',1)[0])
            calls.append(data[0]['id'])
            if data[0]['id']=='s000081' and fail_once[0]:
                fail_once[0]=False
                raise RuntimeError('interrupted')
            return CitedBlock(heading='章节',paragraphs=[{'text':'讲义正文','segment_ids':[data[0]['id']]}])
    monkeypatch.setattr('videosummarizer.learning.OllamaSummarizer',FakeWriter)
    with pytest.raises(Exception): pipeline.run('test')
    assert len(pipeline.state(db.get_job('test'))['blocks'])==1
    pipeline.run('test')
    assert calls.count('s000001')==1
    assert db.get_job('test')['status']=='completed'
    state=pipeline.state(db.get_job('test'))
    result=client.patch('/api/jobs/test/segments/s000081',json={'revision':state['revision'],'text':'用户修订'})
    assert result.status_code==200
    previous=len(calls)
    pipeline.run('test')
    assert calls[previous:]==['s000081']
    assert (directory/'transcript_raw.json').exists()


def test_retry_preserves_files(project):
    directory=create(project)
    project[3].retry('test')
    assert (directory/'transcript_raw.json').exists()


def test_clear_media_requires_saved_transcript(project):
    directory = create(project)
    (directory/'input.mp3').write_bytes(b'only source')
    assert project[-1].delete('/api/jobs/test/media').status_code == 409
    assert (directory/'input.mp3').read_bytes() == b'only source'


def test_enqueue_failure_preserves_import(project, monkeypatch):
    def fail(_):
        raise RuntimeError('queue unavailable')
    monkeypatch.setattr(project[3], 'enqueue', fail)
    content = '1\n00:00:00,000 --> 00:00:02,000\n测试字幕\n'
    assert project[-1].post('/api/import?filename=lesson.srt', content=content.encode()).status_code == 400
    files = list(project[0].jobs_dir.glob('*/input.srt'))
    assert len(files) == 1 and files[0].read_text(encoding='utf-8') == content
    assert project[1].get_job(files[0].parent.name)['status'] == 'failed'


def test_legacy_migration_does_not_forge_citations(project):
    directory=create(project)
    atomic_json(directory/'transcript.json',raw_transcript())
    state=project[2].state(project[1].get_job('test'))
    assert state['legacy'] is True and state['blocks']==[]


def test_transcript_mode_never_initializes_llm_and_exports_content(project, monkeypatch):
    from docx import Document
    from io import BytesIO
    config, db, pipeline, manager, client = project
    def forbidden(*args, **kwargs):
        raise AssertionError('Transcript mode must not call a model')
    monkeypatch.setattr(pipeline.llm_settings, 'runtime', forbidden)
    monkeypatch.setattr('videosummarizer.learning.OllamaSummarizer', forbidden)
    monkeypatch.setattr(pipeline, '_transcribe', forbidden)
    response = client.post('/api/import?filename=lesson.vtt&processing_mode=transcript&llm_provider=openai_compatible',
                           content=b'WEBVTT\n\n00:01.000 --> 00:02.500\nHello &amp; GPU\n')
    assert response.status_code == 202
    job_id = response.json()['id']
    pipeline.run(job_id)
    assert db.get_job(job_id)['status'] == 'completed'
    state = client.get(f'/api/jobs/{job_id}/workspace').json()
    assert state['blocks'] == [] and state['segments'][0]['text'] == 'Hello & GPU'
    for format in ['txt', 'srt', 'json']:
        assert 'Hello & GPU' in client.get(f'/api/jobs/{job_id}/export?format={format}').text
    assert 'Hello &amp; GPU' in client.get(f'/api/jobs/{job_id}/export?format=md').text
    vtt = client.get(f'/api/jobs/{job_id}/export?format=vtt').text
    assert 'WEBVTT' in vtt and '00:00:01.000 --> 00:00:02.500' in vtt and 'Hello &amp; GPU' in vtt
    document = Document(BytesIO(client.get(f'/api/jobs/{job_id}/export?format=docx').content))
    assert any('Hello & GPU' in p.text for p in document.paragraphs)
    assert client.delete(f'/api/jobs/{job_id}/media').status_code == 200
    assert (config.jobs_dir/job_id/'input.vtt').exists()


def test_demo_without_model_and_upgrade_to_lecture(project, monkeypatch):
    config, db, pipeline, manager, client = project
    def forbidden(*args, **kwargs):
        raise AssertionError('Demo must not call a model')
    monkeypatch.setattr('videosummarizer.learning.OllamaSummarizer', forbidden)
    response = client.post('/api/demo')
    assert response.status_code == 201
    job = response.json()
    assert job['status'] == 'completed' and job['processing_mode'] == 'transcript'
    state = pipeline.state(job)
    assert state['segments'] and not state['blocks']
    assert state['needs_proofread'] is False
    response = client.post(f'/api/jobs/{job["id"]}/regenerate', json={'llm_provider':'local'})
    assert response.status_code == 202
    assert db.get_job(job['id'])['processing_mode'] == 'lecture'


def test_asr_transcript_provenance_preserved(project, monkeypatch):
    directory = create(project)
    config, db, pipeline, manager, client = project
    db.update_job('test', processing_mode='transcript')
    data = raw_transcript()
    data['source'] = 'whisper_cpp'
    atomic_json(directory/'transcript_raw.json', data)
    pipeline.run('test')
    assert json.loads((directory/'transcript.json').read_text(encoding='utf-8'))['source'] == 'whisper_cpp'


def test_url_transcript_export_uses_probed_title(project, monkeypatch):
    from docx import Document
    directory = create(project)
    db, pipeline = project[1:3]
    db.update_job('test', processing_mode='transcript')
    def input_transcript(job, cache):
        db.update_job(job['id'], title='New source title')
        return Transcript.model_validate(raw_transcript())
    monkeypatch.setattr(pipeline, '_input_transcript', input_transcript)
    output = pipeline.run('test')
    assert Document(output).paragraphs[0].text == 'New source title'


def test_vtt_export_preserves_multiline_manual_text(project):
    from videosummarizer.text_utils import parse_subtitle
    directory = create(project)
    db, pipeline, manager, client = project[1:]
    pipeline.save(db.get_job('test'), make_workspace(raw_transcript()))
    text = 'First line\n\nSecond line & <literal>\r\n\r\nThird line'
    assert client.patch('/api/jobs/test/segments/s000001', json={'revision':1,'text':text}).status_code == 200
    output = client.get('/api/jobs/test/export?format=vtt')
    path = directory/'roundtrip.vtt'
    path.write_bytes(output.content)
    segments = parse_subtitle(path)
    assert len(segments) == 1 and segments[0].text == 'First line Second line & <literal> Third line'


@pytest.mark.parametrize('ids,expected_calls', [([], 1), (['invented-id'], 2)])
def test_uncertain_citations_remain_explicitly_unverified(project, monkeypatch, ids, expected_calls):
    create(project)
    calls = []
    class FakeWriter:
        def __init__(self,*args): pass
        def ensure_model(self,*args): pass
        def unload_model(self): pass
        def _chat_model(self,*args):
            calls.append(1)
            return CitedBlock(heading='Uncertain', paragraphs=[{'text':'Uncertain claim','segment_ids':ids}])
    monkeypatch.setattr('videosummarizer.learning.OllamaSummarizer', FakeWriter)
    project[2].run('test')
    state = project[2].state(project[1].get_job('test'))
    assert len(calls) == expected_calls
    assert state['blocks'][0]['paragraphs'][0]['segment_ids'] == []
    assert state['blocks'][0]['paragraphs'][0]['review'] == 'unverified'

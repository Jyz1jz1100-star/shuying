from docx import Document

from videosummarizer.article_docx_builder import build_article_docx
from videosummarizer.schemas import ArticleDocument, ArticleSection, SourceMetadata


def test_build_article_docx_has_sections_and_disclosure(tmp_path):
    source = SourceMetadata(
        video_title="示例视频",
        platform="Test",
        author="作者",
        duration_seconds=180,
        source_url="https://example.com/video",
        transcript_language="zh",
        transcription_method="Whisper",
    )
    article = ArticleDocument(
        title="从字幕到完整文章",
        lead="这是一段用于介绍全文的引导文字。",
        sections=[
            ArticleSection(
                heading="第一个主题",
                start_seconds=0,
                end_seconds=180,
                paragraphs=["第一段完整正文。", "第二段继续说明原视频中的信息。"],
            )
        ],
        source=source,
    )
    output = build_article_docx(article, tmp_path)
    document = Document(output)
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    assert headings == ["第一个主题"]
    assert "未分析视频画面" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "完整文章" in output.name

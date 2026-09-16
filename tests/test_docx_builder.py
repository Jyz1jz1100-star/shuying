from docx import Document
from videosummarizer.docx_builder import build_docx
from videosummarizer.schemas import ActionItem, KeyPoint, SourceMetadata, TermDefinition, TimelineItem, VideoSummary


def sample_summary() -> VideoSummary:
    return VideoSummary(
        title="如何把复杂内容讲清楚", one_sentence_summary="视频介绍了从问题定义到结构化表达的一套实用方法。",
        overview="讲者建议先明确听众真正需要解决的问题，再按结论、依据和行动组织内容，并通过复述和反馈检查理解偏差。",
        timeline=[TimelineItem(start_seconds=0,end_seconds=180,heading="定义问题",summary="说明为什么准确的问题定义比立即给答案更重要。"),TimelineItem(start_seconds=180,end_seconds=420,heading="组织表达",summary="使用结论、依据和行动三个层次重组信息。")],
        key_points=[KeyPoint(point="表达前先确认听众和目标。",timestamps=[62]),KeyPoint(point="重要结论需要可验证的依据。",timestamps=[245])],
        action_items=[ActionItem(action="为下一次汇报写出一句话结论。",timestamp=390)],
        terms=[TermDefinition(term="问题定义",definition="在求解前明确目标、边界和判断成功的标准。")], caveats=["示例中的时间安排仅适用于视频展示的情境。"],
        source=SourceMetadata(video_title="如何把复杂内容讲清楚",platform="测试平台",author="测试作者",duration_seconds=420,source_url="https://example.com/video",transcript_language="zh"),
    )


def test_docx_contains_required_structure(tmp_path):
    output = build_docx(sample_summary(), tmp_path); document = Document(output)
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    assert headings == ["摘要","章节时间轴","关键观点","行动项","术语解释","存疑与局限","来源信息"]
    assert len(document.tables) == 1
    assert "_视频总结_" in output.name

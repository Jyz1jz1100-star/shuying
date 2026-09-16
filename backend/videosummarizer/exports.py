from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from .checkpoints import atomic_json
from .evidence import citation, render_markdown
from .text_utils import format_timecode, safe_filename
from .docx_builder import add_hyperlink


def export_workspace(state: dict, job: dict, format: str) -> Path:
    directory = Path(job['job_dir'])
    title = job.get('title') or '述影讲义'
    source = job.get('url') or ''
    path = directory / (safe_filename(title) + '_讲义.' + format)
    if format == 'json':
        atomic_json(path, {**state, 'title': title, 'source_url': source, 'schema_version': 2})
    elif format == 'srt':
        from .learning import transcript_from_state
        from .subtitle_builder import build_srt
        return build_srt(transcript_from_state(state), title, directory)
    elif format == 'md':
        content = render_markdown(state, title, source)
        temporary = path.with_suffix('.md.tmp')
        temporary.write_text(content, encoding='utf-8')
        temporary.replace(path)
    elif format == 'docx':
        document = Document()
        normal = document.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(11)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        document.add_heading(title, 0)
        document.add_paragraph('讲义草稿 · 引用提供原文定位，不代表事实已核实。')
        document.add_paragraph(f'字幕修订版本：{state["revision"]}')
        if source:
            document.add_paragraph('来源：' + source)
        if any(s.get('review') == 'pending' for s in state['segments']):
            document.add_paragraph('包含尚未确认的字幕修改。')
        for block in state['blocks']:
            document.add_heading(block['heading'], 1)
            if block.get('stale'):
                document.add_paragraph('本节依据的字幕已修改，请重新生成。')
            for paragraph in block['paragraphs']:
                document.add_paragraph(paragraph['text'])
                ref = citation(paragraph['segment_ids'], state['segments'], source)
                if not paragraph['segment_ids']:
                    document.add_paragraph('来源待确认')
                else:
                    p = document.add_paragraph()
                    label = f'来源 {format_timecode(ref["start"])}–{format_timecode(ref["end"])}'
                    if ref['url']:
                        add_hyperlink(p, label, ref['url'])
                    else:
                        p.add_run(label)
                    document.add_paragraph(ref['quote'])
        fd, name = tempfile.mkstemp(suffix='.docx', dir=directory)
        os.close(fd)
        temporary = Path(name)
        try:
            document.save(temporary)
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)
    else:
        raise ValueError('不支持的导出格式')
    return path

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .schemas import VideoSummary
from .text_utils import format_timecode, safe_filename


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x17, 0x18, 0x11)
MUTED = RGBColor(0x68, 0x68, 0x60)
GREEN = RGBColor(0x23, 0x4B, 0x3D)
TABLE_FILL = "E8EEF5"
CALLOUT_FILL = "EEF4EC"


def set_run_font(run, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "234B3D")
    borders.append(left)
    p_pr.append(borders)


def add_bottom_rule(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "234B3D")
    borders.append(bottom)
    p_pr.append(borders)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def build_docx(summary: VideoSummary, output_dir: Path) -> Path:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.text = "述影 · 视频内容总结"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], 9, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    set_run_font(run, 9, MUTED)
    add_page_field(footer)
    run = footer.add_run(" 页")
    set_run_font(run, 9, MUTED)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("视频内容总结"), 10, GREEN, bold=True)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run(summary.title), 24, INK, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run(summary.one_sentence_summary), 12.5, MUTED)
    for label, value in (
        ("视频", summary.source.video_title), ("平台", summary.source.platform), ("作者", summary.source.author),
        ("时长", format_timecode(summary.source.duration_seconds)), ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run_font(paragraph.add_run(f"{label}："), 10.5, INK, bold=True)
        set_run_font(paragraph.add_run(value), 10.5, INK)
    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    add_bottom_rule(rule)

    document.add_heading("摘要", level=1)
    overview = document.add_paragraph()
    overview.paragraph_format.left_indent = Inches(0.16)
    overview.paragraph_format.right_indent = Inches(0.12)
    overview.paragraph_format.space_before = Pt(3)
    overview.paragraph_format.space_after = Pt(10)
    shade_paragraph(overview, CALLOUT_FILL)
    set_run_font(overview.add_run(summary.overview), 11, INK)

    document.add_heading("章节时间轴", level=1)
    if summary.timeline:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        headers = ("时间", "章节", "内容摘要")
        for index, text in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), TABLE_FILL)
            cell._tc.get_or_add_tcPr().append(shading)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, 10, DARK_BLUE, bold=True)
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        header_no_split = OxmlElement("w:cantSplit")
        table.rows[0]._tr.get_or_add_trPr().append(header_no_split)
        for item in summary.timeline:
            row = table.add_row()
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            cells = row.cells
            values = (f"{format_timecode(item.start_seconds)}\n至 {format_timecode(item.end_seconds)}", item.heading, item.summary)
            for index, value in enumerate(values):
                cells[index].text = value
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                for run in cells[index].paragraphs[0].runs:
                    set_run_font(run, 9.5 if index == 0 else 10, INK, bold=(index == 1))
        set_table_geometry(table, [1440, 2160, 5760])
    else:
        document.add_paragraph("未识别到足够清晰的章节边界。")

    document.add_heading("关键观点", level=1)
    if summary.key_points:
        for item in summary.key_points:
            suffix = "、".join(format_timecode(value) for value in item.timestamps[:3])
            paragraph = document.add_paragraph(style="List Bullet")
            set_run_font(paragraph.add_run(item.point), 11, INK)
            if suffix:
                set_run_font(paragraph.add_run(f"  [{suffix}]"), 9.5, MUTED)
    else:
        document.add_paragraph("未识别到可独立列出的关键观点。")

    document.add_heading("行动项", level=1)
    if summary.action_items:
        for item in summary.action_items:
            paragraph = document.add_paragraph(style="List Number")
            set_run_font(paragraph.add_run(item.action), 11, INK)
            if item.timestamp is not None:
                set_run_font(paragraph.add_run(f"  [{format_timecode(item.timestamp)}]"), 9.5, MUTED)
    else:
        document.add_paragraph("视频未明确提出行动项。")

    document.add_heading("术语解释", level=1)
    if summary.terms:
        for item in summary.terms:
            paragraph = document.add_paragraph()
            set_run_font(paragraph.add_run(f"{item.term}："), 11, DARK_BLUE, bold=True)
            set_run_font(paragraph.add_run(item.definition), 11, INK)
    else:
        document.add_paragraph("未识别到需要单独解释的术语。")

    document.add_heading("存疑与局限", level=1)
    caveats = [*summary.caveats, "本总结仅分析字幕与语音，未分析视频中的 PPT、图表、代码或其他画面信息。", f"{summary.source.llm_model} 生成内容可能存在遗漏，请以原视频为准。"]
    for caveat in dict.fromkeys(caveats):
        paragraph = document.add_paragraph(style="List Bullet")
        set_run_font(paragraph.add_run(caveat), 10.5, INK)

    document.add_heading("来源信息", level=1)
    source = document.add_paragraph()
    set_run_font(source.add_run("原视频："), 10.5, INK, bold=True)
    add_hyperlink(source, "打开原视频", summary.source.source_url)
    language = document.add_paragraph()
    set_run_font(language.add_run("转写语言："), 10.5, INK, bold=True)
    set_run_font(language.add_run(summary.source.transcript_language), 10.5, INK)
    method = document.add_paragraph()
    set_run_font(method.add_run("内容来源："), 10.5, INK, bold=True)
    set_run_font(method.add_run(summary.source.transcription_method), 10.5, INK)

    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{safe_filename(summary.source.video_title)}_视频总结_{datetime.now():%Y%m%d}.docx"
    destination = output_dir / filename
    document.save(destination)
    validate_docx(destination)
    return destination


def validate_docx(path: Path) -> None:
    document = Document(path)
    headings = {paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")}
    required = {"摘要", "章节时间轴", "关键观点", "行动项", "术语解释", "存疑与局限", "来源信息"}
    missing = required - headings
    if missing:
        raise ValueError(f"DOCX 缺少必要章节：{', '.join(sorted(missing))}")
    if path.stat().st_size < 10_000:
        raise ValueError("DOCX 文件异常或内容不完整")

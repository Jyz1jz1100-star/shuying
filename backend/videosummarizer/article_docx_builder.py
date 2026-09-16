from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .docx_builder import add_hyperlink, add_page_field, set_run_font
from .schemas import ArticleDocument
from .text_utils import format_timecode, safe_filename


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
NAVY = RGBColor(0x20, 0x37, 0x48)
INK = RGBColor(0x1D, 0x22, 0x25)
MUTED = RGBColor(0x68, 0x68, 0x60)
GOLD = RGBColor(0x9A, 0x73, 0x25)


def _configure_narrative_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _add_bottom_rule(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "C9B98E")
    borders.append(bottom)
    properties.append(borders)


def build_article_docx(article: ArticleDocument, output_dir: Path) -> Path:
    document = Document()
    _configure_narrative_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.text = "述影 · 字幕整理完整文章"
    set_run_font(header.runs[0], 9, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("第 "), 9, MUTED)
    add_page_field(footer)
    set_run_font(footer.add_run(" 页"), 9, MUTED)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run(f"完整文章 · 由 {article.source.llm_model} 根据全量字幕整理"), 10, GOLD, bold=True)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run(article.title), 28, NAVY, bold=True)
    lead = document.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.45)
    lead.paragraph_format.right_indent = Inches(0.45)
    lead.paragraph_format.space_after = Pt(24)
    set_run_font(lead.add_run(article.lead), 12.5, MUTED)
    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(30)
    _add_bottom_rule(rule)

    for label, value in (
        ("视频标题", article.source.video_title),
        ("作者", article.source.author),
        ("平台", article.source.platform),
        ("视频时长", format_timecode(article.source.duration_seconds)),
        ("整理时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(f"{label}："), 10, MUTED, bold=True)
        set_run_font(paragraph.add_run(value), 10, MUTED)
    link = document.add_paragraph()
    link.alignment = WD_ALIGN_PARAGRAPH.CENTER
    link.paragraph_format.space_after = Pt(18)
    add_hyperlink(link, "打开原视频", article.source.source_url)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.left_indent = Inches(0.35)
    note.paragraph_format.right_indent = Inches(0.35)
    note.paragraph_format.space_after = Pt(0)
    set_run_font(note.add_run("整理说明："), 9, MUTED, bold=True)
    set_run_font(
        note.add_run(
            f"本文由 {article.source.llm_model} 根据完整字幕和语音转写整理，未分析视频画面。"
            "自动转写与改写可能存在误差；引用、数字和专有名词请以原视频为准。"
        ),
        9,
        MUTED,
    )

    document.add_page_break()
    for index, item in enumerate(article.sections, start=1):
        document.add_heading(item.heading, level=1)
        timecode = document.add_paragraph()
        timecode.paragraph_format.space_after = Pt(8)
        set_run_font(
            timecode.add_run(f"原字幕范围 {format_timecode(item.start_seconds)} - {format_timecode(item.end_seconds)}"),
            9.5,
            MUTED,
            italic=True,
        )
        for paragraph_text in item.paragraphs:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Inches(0.35)
            paragraph.paragraph_format.widow_control = True
            set_run_font(paragraph.add_run(paragraph_text), 11, INK)

    output_dir.mkdir(parents=True, exist_ok=True)
    destination = output_dir / f"{safe_filename(article.source.video_title)}_完整文章_{datetime.now():%Y%m%d}.docx"
    document.save(destination)
    _validate_article_docx(destination, len(article.sections))
    return destination


def _validate_article_docx(path: Path, expected_sections: int) -> None:
    document = Document(path)
    headings = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    if len(headings) < expected_sections:
        raise ValueError("完整文章 DOCX 缺少正文章节")
    if "未分析视频画面" not in "\n".join(paragraph.text for paragraph in document.paragraphs):
        raise ValueError("完整文章 DOCX 缺少画面分析范围说明")
    if path.stat().st_size < 10_000:
        raise ValueError("完整文章 DOCX 文件异常或内容不完整")
